from pathlib import Path
import re, json, html, sys
PREVIEW = "--preview" in sys.argv
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
report = ROOT / 'docs/report/软件课程设计报告-刘东昕.md'
text = report.read_text()

def excerpt(path, start, end, limit=48):
    src = (ROOT/path).read_text()
    begin = src.index(start)
    finish = src.index(end, begin)
    block = src[begin:finish].strip().splitlines()
    return '```rust\n' + '\n'.join(block[:limit]) + '\n```'

extract = (ROOT/'src-tauri/src/ai/extraction.rs').read_text()
prompt = extract[extract.index('pub(super) const PROMPT:'):extract.index('\nfn key(')].strip()
replacements = {
    '{{EXTRACTION_PROMPT}}': '```rust\n'+prompt+'\n```',
    '{{RESPONSE_RECOVERY}}': excerpt('src-tauri/src/ai/extraction.rs', 'fn response_points(', '// Validate complete returned', 60),
    '{{RETRIEVAL_CODE}}': excerpt('src-tauri/src/qa.rs', 'fn terms(', '    if let Some(k)', 48),
    '{{IPC_CODE}}': excerpt('src-tauri/src/lib.rs', '    #[tauri::command]\n    async fn organize(', '    #[tauri::command]\n    async fn rebuild_graph', 40),
}
unitlog = (ROOT/'src-tauri/target/guiye-unit-tests.log').read_text()
unit_match = re.search(r'test result: ok\. (\d+) passed; 0 failed',unitlog)
ui_log = (ROOT/'src-tauri/target/guiye-browser-tests-final.log')
ui_match = re.search(r'(\d+) passed',ui_log.read_text()) if ui_log.exists() else None
freshlog = (ROOT/'src-tauri/target/guiye-general-fresh/test.log').read_text()
fresh_match = re.search(r'LIVE EXTRACTION PASS: (\d+) points',freshlog)
usage = re.search(r'LIVE USAGE: (\d+) requests, (\d+) input tokens, (\d+) output tokens',freshlog)
full_log_path = ROOT/'src-tauri/target/guiye-general-full/test.log'
full_log = full_log_path.read_text()
full = re.search(r'LIVE ORGANIZE PASS: (\d+) points, (\d+) nodes, (\d+) notes, depth (\d+)',full_log)
full_usage = re.search(r'LIVE USAGE: (\d+) requests, (\d+) input tokens, (\d+) output tokens',full_log)
full_passed = bool(full and full_usage)
if not full_passed:
    full = ['unverified', '0', '0', '0', '0']
    full_usage = ['unverified', '0', '0', '0']
if not all([unit_match,ui_match,fresh_match,usage,full,full_usage]):
    raise SystemExit('等待最终验收记录齐全，不生成带未完成测试结论的提交版。')
replacements['{{TEST_RESULTS}}'] = f'''| 测试层次 | 实际结果 | 记录位置 |
| --- | --- | --- |
| Rust 回归测试 | {unit_match[1]} 项通过，0 项失败 | guiye-unit-tests.log |
| Playwright 界面测试 | {ui_match[1]} 项通过 | guiye-browser-tests-final.log |
| 实际模型从零提取 | 5 份笔记，{fresh_match[1]} 个知识点，未使用原文兜底 | guiye-general-fresh/test.log |
| 实际模型完整整理 | {full[1]} 个知识点、{full[2]} 个图谱节点、{full[3]} 篇文档 | guiye-general-full/test.log |'''
replacements['{{LIVE_RESULTS}}'] = f'''2026年9月8日，使用 ~/note/OS 中的5份课程笔记，通过用户配置的 deepseek-v4-pro 完成实际服务测试。授权确认后运行测试，测试产物写入构建目录，不以测试脚本覆盖用户工作区。

从零提取共调用模型 {usage[1]} 次，得到 {fresh_match[1]} 个知识点；实际输入 {int(usage[2]):,} tokens，输出 {int(usage[3]):,} tokens。所有来源均有知识点关联，未出现原文兜底条目，也未要求用户中途点击继续。

完整整理测试在已有项目基础上更新，得到 {full[1]} 个知识点、{full[2]} 个图谱节点与 {full[3]} 篇文档。图谱最大深度为 {full[4]} 条父子边，即包含根节点在内的 {int(full[4])+1} 层；全部节点及文档通过引用和结构校验，未保留提取待继续或层级待优化标记。该测试调用 {full_usage[1]} 次模型，输入 {int(full_usage[2]):,} tokens，输出 {int(full_usage[3]):,} tokens。两次测试的知识点数量不同，是因为完整测试复用了已有项目内容并执行去重与纠错，不能把两者数量差直接解释为覆盖率。

以上数据来自实际服务返回的 usage 字段及测试断言，不是估算的费用或模拟服务结果。与修复前相比，已确认消除了同一批无进展后反复要求用户继续的流程。核验请求也移除了重复原文证据，但模型输出长度存在变化，不能仅凭一次运行承诺固定的 token 节省比例。'''
if not full_passed:
    replacements['{{TEST_RESULTS}}'] = replacements['{{TEST_RESULTS}}'].replace('| 实际模型完整整理 | 0 个知识点、0 个图谱节点、0 篇文档 | guiye-general-full/test.log |', '| 旧数据完整整理 | 前次发现旧原文条目被复用；针对性回归已通过，完整实测尚未复验 | guiye-general-full/test.log |')
    replacements['{{LIVE_RESULTS}}'] = f'''2026年9月8日，使用 ~/note/OS 中的5份课程笔记，通过配置的 deepseek-v4-pro 完成从零提取的实际服务测试。测试产物写入构建目录，测试脚本不覆盖用户工作区。

从零提取共调用模型 {usage[1]} 次，得到 {fresh_match[1]} 个知识点；服务返回的实际输入为 {int(usage[2]):,} tokens，输出为 {int(usage[3]):,} tokens。五份来源均有知识点关联，未出现原文兜底条目，也未要求中途点击继续。本次结果证明了实际笔记与通用提取协议的兼容性，但不等价于对全部语义内容的人工核验。

旧项目的完整整理测试暴露了另一项兼容性问题：旧版本标记为原文保留的两条记录被直接复用，使流程虽然完成文档写作，最终仍未满足“无原文兜底”的验收断言。随后增加了针对旧记录的上下文复核，并通过确定性回归验证：仅重处理原文条目，有效知识点保持原 ID，处理结束不返回待继续状态。

截至本报告记录时，这一兼容性修复的真实服务完整复测尚未完成。因此，报告不将前次完整整理写为通过，也不报告未经核验的层级深度、文档数量或全流程 token 节省比例。成本优化的实现包括复用有效缓存、截断后只补缺失段落、缩短重复上下文，以及以两个小节为上限的并发写作；优化幅度仍需在相同输入和设置下重复测量。'''
if PREVIEW:
    replacements['{{LIVE_RESULTS}}'] = '排版预览：最终完整流程复测尚未结束，本页不作验收结论。'
for key,value in replacements.items(): text=text.replace(key,value)
assert not re.search(r'\{\{[A-Z_]+\}\}',text)
if not PREVIEW:
    (ROOT/'docs/report/最终报告正文.md').write_text(text)

doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(2.5);section.bottom_margin=Cm(2.5)
section.left_margin=Cm(2.8);section.right_margin=Cm(2.6)
section.header_distance=Cm(1.3);section.footer_distance=Cm(1.3)
styles=doc.styles
for style_name in ['Normal','Body Text','Caption']:
    st=styles[style_name];st.font.name='Times New Roman';st.font.size=Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    st.paragraph_format.line_spacing=1.5
    st.paragraph_format.space_after=Pt(5)
for level in [1,2,3]:
    st=styles[f'Heading {level}'];st.font.name='Noto Sans CJK SC';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt({1:16,2:14,3:12}[level]);st.font.bold=True
    st.font.italic=False;st.font.complex_script=False
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        st.element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
    for attr in ['ascii','hAnsi','eastAsia','cs']:
        st.element.rPr.rFonts.set(qn('w:'+attr),'Noto Sans CJK SC')
    for tag in ['i','iCs']:
        el=st.element.rPr.find(qn('w:'+tag))
        if el is None:el=OxmlElement('w:'+tag);st.element.rPr.append(el)
        el.set(qn('w:val'),'0')
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(8)
styles['Heading 1'].paragraph_format.page_break_before=True
if 'Source Code' not in styles:
    from docx.enum.style import WD_STYLE_TYPE
    st=styles.add_style('Source Code',WD_STYLE_TYPE.PARAGRAPH)
else: st=styles['Source Code']
st.font.name='Consolas';st.font.size=Pt(8.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
st.font.italic=False
st.paragraph_format.line_spacing=1.05;st.paragraph_format.space_after=Pt(0)
st.paragraph_format.left_indent=Cm(.2)

# Cover follows the school template; grading and signature fields remain blank.
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('软件课程设计报告');r.bold=True;r.font.size=Pt(24)
r.font.name='黑体';r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体')
p=doc.add_paragraph('(2026年春季学期)');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('序号：');p.paragraph_format.space_before=Pt(18)
for title in ['题目：归页——基于大语言模型的','本地知识整理系统']:
    p=doc.add_paragraph(title);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.bold=True;r.font.size=Pt(18)
doc.add_paragraph()
for line in ['系别：计算机科学与技术','班级：计实验23','姓名：刘东昕','学号：23101020202']:
    p=doc.add_paragraph(line);p.paragraph_format.left_indent=Cm(2);p.paragraph_format.space_after=Pt(8)
doc.add_paragraph('总成绩：')
p=doc.add_paragraph('评语：');p.paragraph_format.space_after=Pt(65)
doc.add_paragraph('指导教师签字：________________    日期：________________')
doc.add_page_break()
p=doc.add_paragraph('目录');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold=True;p.runs[0].font.size=Pt(18)
p=doc.add_paragraph()
begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-2" \\h \\z \\u '
separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
for el in [begin,instruction,separate]:p.add_run()._r.append(el)
# Cached contents remain readable even before a word processor updates fields.
for line in text.splitlines():
    if line.startswith('## '):p.add_run(line[3:]+'\n')
p.add_run()._r.append(end)
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)

lines=text[text.index('## 1 '):].splitlines()
i=0
while i<len(lines):
    line=lines[i]
    if not line.strip():i+=1;continue
    if line.startswith('!['):
        match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
        caption,filename=match.groups()
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True
        width=12.4 if 'organize-flow' in filename else 15.5
        p.add_run().add_picture(str(report.parent/filename),width=Cm(width))
        p=doc.add_paragraph(caption,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=False
        for r in p.runs:r.font.size=Pt(10);r.font.italic=False
        i+=1;continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=doc.add_paragraph(lines[i],style='Source Code')
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F4F4F4');p._p.get_or_add_pPr().append(shading)
            i+=1
        doc.add_paragraph();i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[c.strip() for c in lines[i].strip('|').split('|')]
            if not all(re.match(r'^:?-+:?$',c) for c in cells):rows.append(cells)
            i+=1
        t=doc.add_table(rows=1,cols=len(rows[0]));t.style='Table Grid'
        for j,c in enumerate(rows[0]):t.rows[0].cells[j].text=c
        for row in rows[1:]:
            cells=t.add_row().cells
            for j,c in enumerate(row):cells[j].text=c
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                for run in p.runs:run.bold=True
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.15
                    for run in p.runs:run.font.size=Pt(10)
        doc.add_paragraph();continue
    if line.startswith('### '):doc.add_heading(line[4:],2)
    elif line.startswith('## '):doc.add_heading(line[3:],1)
    else:
        p=doc.add_paragraph(line)
        p.paragraph_format.first_line_indent=Cm(.74)
    i+=1
section.different_first_page_header_footer=True
header=section.header.paragraphs[0];header.text='软件课程设计报告 · 归页';header.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in header.runs:r.font.size=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('— ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
footer.add_run(' —')
doc.core_properties.title='归页：基于大语言模型的本地知识整理系统'
doc.core_properties.author='刘东昕'
doc.core_properties.subject='软件课程设计报告'
output = ROOT / ('src-tauri/target/report-preview/报告排版预览.docx' if PREVIEW else '软件课程设计报告-刘东昕-提交版.docx')
doc.save(output)
loaded=Document(output)
assert len(loaded.tables)>=2
assert len(loaded.inline_shapes)==7
for level in [1,2,3]:
    assert loaded.styles[f'Heading {level}'].font.italic is False
assert len([p for p in loaded.paragraphs if p.style.name=='Heading 1'])==6
assert all(v in '\n'.join(p.text for p in loaded.paragraphs) for v in ['刘东昕','23101020202','计实验23'])
print(f'Word report saved: {output}; {len(loaded.paragraphs)} paragraphs; {len(loaded.tables)} tables')
